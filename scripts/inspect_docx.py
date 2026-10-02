from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    path = Path(sys.argv[1])
    document = Document(path)
    payload = {
        "paragraphs": [
            {
                "index": index,
                "text": paragraph.text,
                "runs": [run.text for run in paragraph.runs],
            }
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip()
        ],
        "tables": [
            [
                [cell.text for cell in row.cells]
                for row in table.rows
            ]
            for table in document.tables
        ],
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
